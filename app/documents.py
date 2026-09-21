"""Documents on the stick (D-79): the harness reads them, the model reads the extracted text.

A 1B model does not "see" a file. What it can do is read text, so the harness extracts text from
the document and either hands the model the first page (`read_file`) or indexes the whole thing
beside the packs so `pack_search` can find the paragraph that answers (`/index`). That is the
same mechanism the shelf manuals use, and it works at any model size; what a bigger model buys is
a longer window, not a new sense.

Formats: .txt .md .csv .json .log (read as text); .pdf (pypdf, BSD-3: the text layer only — a
scanned PDF has none and the reader says so; OCR is not on this stick); .docx (python-docx, MIT:
paragraphs and tables). Anything else: refused by name. Images: not in this version.
"""

from __future__ import annotations

import io
import os
import re

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".rst"}


def extract_text(path: str, max_chars: int = 2_000_000) -> tuple[str, str]:
    """(text, note). note says what was done or why nothing came out."""
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXT:
        t = io.open(path, encoding="utf-8", errors="replace").read(max_chars)
        return t, f"plain text, {len(t):,} chars"
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return "", "PDF support is not built into this program"
        try:
            r = PdfReader(path)
        except Exception as e:  # noqa: BLE001
            return "", f"could not open the PDF: {str(e)[:80]}"
        pages = []
        for i, pg in enumerate(r.pages):
            try:
                pages.append(pg.extract_text() or "")
            except Exception:  # noqa: BLE001
                pages.append("")
            if sum(len(p) for p in pages) > max_chars:
                break
        text = "\n\n".join(f"<!-- page {i + 1} -->\n{p.strip()}" for i, p in enumerate(pages) if p.strip())
        if len(re.findall(r"[A-Za-z]{3,}", text)) < max(5, 3 * len(r.pages)):   # a scan yields nothing; a real page yields words
            return "", f"the PDF has {len(r.pages)} pages but no usable text layer (a scan?); this program has no OCR"
        return text, f"PDF, {len(r.pages)} pages, {len(text):,} chars of text layer"
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            return "", "Word support is not built into this program"
        try:
            d = docx.Document(path)
        except Exception as e:  # noqa: BLE001
            return "", f"could not open the document: {str(e)[:80]}"
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n\n".join(parts)[:max_chars]
        return text, f"Word document, {len(d.paragraphs)} paragraphs, {len(d.tables)} tables, {len(text):,} chars"
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff"):
        return "", "an image: this version reads text, not pictures (no OCR on the stick)"
    return "", f"unsupported type {ext or '(none)'}: text, PDF and Word documents only"


def index_into(path: str, cache_dir: str) -> tuple[str, str]:
    """Extract a document (or every document in a folder) into cache_dir as .txt sidecars the pack
    index can read. Returns (summary, first sidecar path)."""
    os.makedirs(cache_dir, exist_ok=True)
    files = []
    if os.path.isdir(path):
        for dp, dns, fns in os.walk(path):
            dns[:] = [d for d in dns if not d.startswith(".")]          # never walk the .text cache itself
            for fn in fns:
                if os.path.splitext(fn)[1].lower() in TEXT_EXT | {".pdf", ".docx"}:
                    files.append(os.path.join(dp, fn))
    else:
        files.append(path)
    done, notes, first = 0, [], ""
    for f in files:
        text, note = extract_text(f)
        base = re.sub(r"[^A-Za-z0-9_.\-]+", "_", os.path.basename(f))
        if os.path.abspath(os.path.dirname(f)) == os.path.abspath(cache_dir):
            continue                                                     # a sidecar is not a document
        if not text:
            notes.append(f"{os.path.basename(f)}: {note}"); continue
        out = os.path.join(cache_dir, base + ".txt")
        io.open(out, "w", encoding="utf-8", newline="\n").write(f"[document: {os.path.basename(f)} — {note}]\n\n" + text)
        first = first or out
        done += 1
    summary = f"indexed {done} of {len(files)} document(s) into {cache_dir}"
    if notes:
        summary += "\n  skipped: " + "; ".join(notes[:6])
    return summary, first
